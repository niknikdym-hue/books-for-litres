#!/usr/bin/env python3
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Cm, Pt
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[3]
BOOK = ROOT / "sekrety-prodvizheniya-uslug" / "kak-prodavat-uslugi"
MANUSCRIPT = BOOK / "manuscript"
NOTES = BOOK / "reader-notes" / "NOTES-AND-SOURCES.md"
OUT = BOOK / "deliverables"
OUT.mkdir(parents=True, exist_ok=True)

FILES = [
    MANUSCRIPT / "00-vvedenie.md",
    MANUSCRIPT / "01-pochemu-horoshuyu-uslugu-trudno-kupit.md",
    MANUSCRIPT / "02-chto-klient-dolzhen-reshitsya-kupit.md",
    MANUSCRIPT / "03-dokazatelstva-vmesto-uvereniy.md",
    MANUSCRIPT / "04-pokazat-rabotu-do-nachala-raboty.md",
    MANUSCRIPT / "05-diagnoz-do-predlozheniya.md",
    MANUSCRIPT / "06-kogda-horoshaya-prodazha-zakanchivaetsya-otkazom.md",
    MANUSCRIPT / "07-kommercheskoe-predlozhenie-kak-dokument-resheniya.md",
    MANUSCRIPT / "08-tsena-obem-i-risk.md",
    MANUSCRIPT / "09-ya-podumayu-eto-ne-odno-vozrazhenie.md",
    MANUSCRIPT / "10-pochemu-vy-proigrali-sdelku.md",
]

DOCX = OUT / "kak-prodavat-uslugi-current-manuscript.docx"
TXT = OUT / "kak-prodavat-uslugi-current-manuscript.txt"

for p in FILES + [NOTES]:
    if not p.exists():
        raise SystemExit(f"Missing required reader file: {p}")


def strip_md_for_txt(text: str) -> str:
    text = re.sub(r"^#{1,6}\s+", "", text, flags=re.M)
    text = text.replace("**", "").replace("`", "")
    text = re.sub(r"^>\s?", "", text, flags=re.M)
    return text.strip()


def add_inline(paragraph, text: str):
    # Minimal renderer for bold and inline code; no style markup remains visible.
    token = re.compile(r"(\*\*.*?\*\*|`.*?`)")
    pos = 0
    for m in token.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        t = m.group(0)
        if t.startswith("**"):
            r = paragraph.add_run(t[2:-2])
            r.bold = True
        else:
            r = paragraph.add_run(t[1:-1])
            r.font.name = "Aptos Mono"
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_markdown(doc: Document, text: str):
    lines = text.splitlines()
    for raw in lines:
        line = raw.rstrip()
        if not line:
            doc.add_paragraph()
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, line[2:])
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, line[3:])
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 3")
            add_inline(p, line[4:])
            continue
        if re.match(r"^[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^[-*]\s+", "", line))
            continue
        if re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s+", "", line))
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.right_indent = Cm(0.4)
            add_inline(p, line[2:])
            continue
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(3)
        add_inline(p, line)


# TXT
parts = [
    "КАК ПРОДАВАТЬ УСЛУГИ\n",
    "Серия «Секреты продвижения услуг», книга №1\n",
    "Текущая полная рукопись. Версия 2026-09-07.\n",
    "\n" + ("=" * 68) + "\n",
]
for i, p in enumerate(FILES):
    parts.append(strip_md_for_txt(p.read_text(encoding="utf-8")))
    parts.append("\n\n" + ("=" * 68) + "\n\n")
parts.append(strip_md_for_txt(NOTES.read_text(encoding="utf-8")))
TXT.write_text("\n".join(parts).strip() + "\n", encoding="utf-8")

# DOCX
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.2)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.0)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
normal.font.size = Pt(11.5)

for style_name, size in [("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)]:
    st = styles[style_name]
    st.font.name = "Aptos"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    st.font.size = Pt(size)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("КАК ПРОДАВАТЬ УСЛУГИ")
r.bold = True
r.font.size = Pt(24)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Серия «Секреты продвижения услуг», книга №1")
r.font.size = Pt(13)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Текущая полная рукопись · 7 сентября 2026")
r.italic = True
r.font.size = Pt(10.5)

doc.add_page_break()

for idx, pth in enumerate(FILES):
    if idx:
        doc.add_page_break()
    add_markdown(doc, pth.read_text(encoding="utf-8"))

doc.add_page_break()
add_markdown(doc, NOTES.read_text(encoding="utf-8"))

doc.core_properties.title = "Как продавать услуги"
doc.core_properties.subject = "Серия «Секреты продвижения услуг», книга №1"
doc.core_properties.author = ""
doc.core_properties.comments = "Current complete manuscript generated from repository source of truth."
doc.save(DOCX)

print(DOCX)
print(TXT)
