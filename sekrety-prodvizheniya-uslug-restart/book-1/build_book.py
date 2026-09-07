"""Build the complete reading manuscript from the canonical chapter files.

Run with the bundled primary-runtime Python; python-docx is required.
The DOCX uses narrative_proposal with the named book_reading overrides below.
"""
from pathlib import Path
import argparse, json, re, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

BASE=Path(__file__).resolve().parent
PARTS={1:'Часть I. Выбрать задачу и определить результат',4:'Часть II. Сформировать предложение',8:'Часть III. Согласовать покупку',12:'Часть IV. Начать работу и проверить изменения'}
TOKENS={
 'preset':'narrative_proposal','header_pattern':'editorial_cover',
 'page':{'width_in':8.5,'height_in':11,'margin_in':1,'header_footer_in':.492,'width_dxa':9360},
 'book_reading':{'font':'Times New Roman','size_pt':12,'line':1.15,'before_pt':0,'after_pt':3,'alignment':'justified','heading_color':'202020'},
 'h1':{'size_pt':18,'before_pt':18,'after_pt':10,'line':1.15},
 'h2':{'size_pt':13,'before_pt':12,'after_pt':6,'line':1.15},
 'h3':{'size_pt':12,'before_pt':8,'after_pt':4,'line':1.15},
 'table':{'font':'Arial','size_pt':10.5,'line':1.1,'before_pt':0,'after_pt':0,'width_dxa':9360,'indent_dxa':120,'margins_dxa':[80,80,120,120],'fill':'F4F6F9','border':'CCCCCC'},
 'book_excerpt':{'size_pt':11.5,'line':1.15,'before_pt':4,'after_pt':5,'indent_dxa':240,'border':'BBBBBB'},
 'bibliography_list':{'size_pt':10.5,'line':1.1,'before_pt':0,'after_pt':7,'text_dxa':540,'marker_dxa':300,'hanging_dxa':240,'marker_alignment':'right'},
 'toc':{'size_pt':11,'line':1.05,'before_pt':0,'after_pt':4},
 'furniture':{'font':'Arial','size_pt':8.5,'color':'777777','before_pt':0,'after_pt':0,'line':1},
 'table_citation_text':{'before_pt':4,'after_pt':4},
 'overrides_reason':'Serif typography and restrained monochrome hierarchy for a full-length book; compact paragraph rhythm keeps a readable manuscript rather than proposal spacing.'
}

def xml(tag,**attrs):
 e=OxmlElement('w:'+tag)
 for k,v in attrs.items():e.set(qn('w:'+k),str(v))
 return e

def style(doc,name,size=12,before=0,after=4,line=1.15,font='Times New Roman',bold=False,color='202020',align=WD_ALIGN_PARAGRAPH.LEFT):
 s=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 s.font.name=font;s.font.size=Pt(size);s.font.bold=bold;s.font.italic=False;s.font.color.rgb=RGBColor.from_string(color)
 for e in list(s.element.get_or_add_pPr().findall(qn('w:pBdr'))):s.element.get_or_add_pPr().remove(e)
 r=s.element.get_or_add_rPr();rf=r.find(qn('w:rFonts'))
 if rf is None:rf=xml('rFonts');r.insert(0,rf)
 for a in ('ascii','hAnsi','eastAsia','cs'):rf.set(qn('w:'+a),font)
 for a in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):rf.attrib.pop(qn('w:'+a),None)
 for tag in ('spacing','kern'):
  for e in list(r.findall(qn('w:'+tag))):r.remove(e)
 lang=r.find(qn('w:lang'))
 if lang is None:lang=xml('lang');r.append(lang)
 lang.set(qn('w:val'),'ru-RU')
 pf=s.paragraph_format;pf.space_before=Pt(before);pf.space_after=Pt(after);pf.line_spacing=line;pf.alignment=align
 pf.left_indent=Pt(0);pf.right_indent=Pt(0);pf.first_line_indent=Pt(0);pf.widow_control=True;pf.keep_with_next=False;pf.keep_together=False
 return s

def bookmark(p,name,num):
 p._p.append(xml('bookmarkStart',id=num,name=name));p._p.append(xml('bookmarkEnd',id=num))

def link(p,label,url=None,anchor=None):
 h=xml('hyperlink')
 if url:h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 else:h.set(qn('w:anchor'),anchor);h.set(qn('w:history'),'1')
 r=xml('r');rp=xml('rPr');rp.append(xml('rStyle',val='Hyperlink'));r.append(rp);t=xml('t');t.text=label;r.append(t);h.append(r);p._p.append(h)

TOKEN=re.compile(r'(\[[^\]]+\]\([^)]+\)|\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)')
def plain(s):
 s=re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',s)
 return s.replace('**','').replace('*','').replace('`','')

def inline(p,s):
 for bit in TOKEN.split(s):
  if not bit:continue
  m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',bit)
  if m:link(p,m.group(1),url=m.group(2));continue
  if bit.startswith('**') and bit.endswith('**'):p.add_run(bit[2:-2]).bold=True
  elif bit.startswith('*') and bit.endswith('*'):p.add_run(bit[1:-1]).italic=True
  elif bit.startswith('`') and bit.endswith('`'):p.add_run(bit[1:-1])
  else:p.add_run(bit)
 actual=''.join(p._p.itertext()) if False else ''.join(e.text or '' for e in p._p.iter(qn('w:t')))
 assert actual==plain(s),(actual,plain(s))

def num_id(doc):
 n=doc.part.numbering_part.element
 aid=max([int(e.get(qn('w:abstractNumId'))) for e in n.findall(qn('w:abstractNum'))]+[0])+1
 nid=max([int(e.get(qn('w:numId'))) for e in n.findall(qn('w:num'))]+[0])+1
 a=xml('abstractNum',abstractNumId=aid);a.append(xml('multiLevelType',val='singleLevel'))
 lvl=xml('lvl',ilvl=0);lvl.append(xml('start',val=1));lvl.append(xml('numFmt',val='decimal'));lvl.append(xml('lvlText',val='%1.'));lvl.append(xml('lvlJc',val='right'));lvl.append(xml('suff',val='tab'))
 pp=xml('pPr');tabs=xml('tabs');tabs.append(xml('tab',val='num',pos=540));pp.append(tabs);pp.append(xml('ind',left=540,hanging=240));lvl.append(pp);a.append(lvl);n.append(a)
 item=xml('num',numId=nid);item.append(xml('abstractNumId',val=aid));n.append(item);return nid

def widths(ch,idx,n):
 specific={(1,0):[2500,6860],(2,0):[2280,3540,3540],(4,0):[2600,6760],(5,0):[6800,2560],(5,1):[4200,2580,2580],(6,0):[1250,2290,1940,1940,1940],(7,0):[2600,2500,2280,1980],(8,0):[2900,3300,3160],(9,0):[2100,3350,3910],(10,0):[2600,2500,1460,2800],(11,0):[2000,3800,3560],(12,0):[2100,2000,2800,2460],(13,0):[2000,3800,3560],(13,1):[1900,7460],(14,0):[5360,2000,2000],(14,1):[2400,6960]}
 w=specific[(ch,idx)];assert len(w)==n and sum(w)==9360;return w

def add_table(doc,rows,ch,idx,records):
 n=len(rows[0]);assert all(len(r)==n for r in rows)
 w=widths(ch,idx,n);table=doc.add_table(rows=len(rows),cols=n);table.autofit=False
 pr=table._tbl.tblPr
 for tag in ['tblW','tblInd','tblLayout','tblCellMar','tblBorders']:
  for e in list(pr.findall(qn('w:'+tag))):pr.remove(e)
 pr.append(xml('tblW',w=9360,type='dxa'));pr.append(xml('tblInd',w=120,type='dxa'));pr.append(xml('tblLayout',type='fixed'))
 mar=xml('tblCellMar')
 for k,v in [('top',80),('bottom',80),('start',120),('end',120)]:mar.append(xml(k,w=v,type='dxa'))
 pr.append(mar);b=xml('tblBorders')
 for k in ['top','left','bottom','right','insideH','insideV']:b.append(xml(k,val='single',sz=4,color='CCCCCC'))
 pr.append(b);grid=table._tbl.tblGrid
 for e in list(grid):grid.remove(e)
 for cw in w:grid.append(xml('gridCol',w=cw))
 for ri,row in enumerate(rows):
  trpr=table.rows[ri]._tr.get_or_add_trPr();trpr.append(xml('cantSplit'))
  if ri==0:trpr.append(xml('tblHeader'))
  for ci,text in enumerate(row):
   c=table.cell(ri,ci);tcpr=c._tc.get_or_add_tcPr();tw=tcpr.find(qn('w:tcW'));tw.set(qn('w:w'),str(w[ci]));tw.set(qn('w:type'),'dxa')
   if ri==0:tcpr.append(xml('shd',fill='F4F6F9',val='clear'))
   p=c.paragraphs[0];p.style='Table Header Text' if ri==0 else ('Table Text Narrow' if n==5 else 'Table Text');inline(p,text)
   records.append({'kind':'table_cell','chapter':ch,'text':plain(text)})
 doc.add_paragraph(style='Table After')

def body(doc,text,ch,records,bibnum):
 lines=text.splitlines();i=0;ti=0
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;continue
  if line.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].strip().startswith('|'):
    cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
    if not all(re.fullmatch(r':?-+:?',x) for x in cells):rows.append(cells)
    i+=1
   add_table(doc,rows,ch,ti,records);ti+=1;continue
  h=re.match(r'^(#{1,3}) (.+)',line)
  if h:
   lvl=len(h.group(1));title=h.group(2)
   if lvl==1 and ch in PARTS:
    p=doc.add_paragraph(PARTS[ch],style='Part Label');p.paragraph_format.page_break_before=True
   p=doc.add_paragraph(style='Heading '+str(lvl));inline(p,title)
   if lvl==1:
    p.paragraph_format.page_break_before=(ch not in PARTS);bookmark(p,'chapter_'+str(ch),ch+20)
   records.append({'kind':'heading','chapter':ch,'text':plain(title)})
   i+=1;continue
  quoted=line.startswith('>');parts=[]
  while i<len(lines):
   line=lines[i].strip()
   if not line or line.startswith(('#','|')) or (line.startswith('>')!=quoted):break
   if quoted:
    line=re.sub(r'^>\s?','',line)
    if not line:i+=1;break
   parts.append(line);i+=1
  if not parts:continue
  s=' '.join(parts);is_bib=ch==99 and re.match(r'^\d+\. ',s)
  if is_bib:s=re.sub(r'^\d+\. ','',s)
  p=doc.add_paragraph(style='Bibliography Text' if is_bib else 'Book Excerpt' if quoted else 'Normal');inline(p,s)
  if is_bib:
   pp=p._p.get_or_add_pPr();np=xml('numPr');np.append(xml('ilvl',val=0));np.append(xml('numId',val=bibnum));pp.append(np)
  records.append({'kind':'bibliography' if is_bib else 'quote' if quoted else 'paragraph','chapter':ch,'text':plain(s)})

def make_doc(files,out):
 doc=Document();sec=doc.sections[0]
 sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 for k in ('top_margin','bottom_margin','left_margin','right_margin'):setattr(sec,k,Inches(1))
 sec.header_distance=Inches(.492);sec.footer_distance=Inches(.492);sec.different_first_page_header_footer=True
 style(doc,'Normal',after=3,align=WD_ALIGN_PARAGRAPH.JUSTIFY)
 for n in (1,2,3):
  t=TOKENS['h'+str(n)];s=style(doc,'Heading '+str(n),t['size_pt'],t['before_pt'],t['after_pt'],bold=True);s.paragraph_format.keep_with_next=True;s.paragraph_format.keep_together=True
 style(doc,'Title',32,0,18,1.06,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER)
 style(doc,'Subtitle',11,0,18,1.1,align=WD_ALIGN_PARAGRAPH.CENTER,color='666666')
 style(doc,'Cover Series',10.5,96,22,1.1,font='Arial',align=WD_ALIGN_PARAGRAPH.CENTER,color='666666')
 style(doc,'Cover Author',18,0,26,1.1,align=WD_ALIGN_PARAGRAPH.CENTER)
 style(doc,'Cover Year',11,96,0,1.1,align=WD_ALIGN_PARAGRAPH.CENTER,color='666666')
 s=style(doc,'Contents Title',18,0,16,bold=True);s.paragraph_format.page_break_before=True;s.paragraph_format.keep_with_next=True
 style(doc,'Contents Entry',11,0,4,1.05)
 s=style(doc,'Contents Part',10,10,5,1.05,font='Arial',bold=True,color='666666');s.paragraph_format.keep_with_next=True
 s=style(doc,'Part Label',10,0,8,1.1,font='Arial',bold=True,color='666666');s.paragraph_format.keep_with_next=True
 style(doc,'Table Text',10.5,0,0,1.1,font='Arial')
 style(doc,'Table Text Narrow',10,0,0,1.1,font='Arial')
 style(doc,'Table Header Text',10,0,0,1.1,font='Arial',bold=True)
 style(doc,'Table After',1,0,4,1)
 s=style(doc,'Book Excerpt',11.5,4,5,1.15);s.paragraph_format.left_indent=Pt(12);s.paragraph_format.right_indent=Pt(12);s.paragraph_format.keep_together=True
 pp=s.element.get_or_add_pPr();bd=xml('pBdr');bd.append(xml('left',val='single',sz=6,space=8,color='BBBBBB'));pp.append(bd)
 s=style(doc,'Bibliography Text',10.5,0,7,1.1);s.paragraph_format.left_indent=Pt(27);s.paragraph_format.first_line_indent=Pt(-12)
 style(doc,'Table Citation',10.5,4,4,1.1)
 for name in ['Header','Footer']:style(doc,name,8.5,0,0,1,font='Arial',color='777777')
 hs=doc.styles['Hyperlink'] if 'Hyperlink' in doc.styles else doc.styles.add_style('Hyperlink',WD_STYLE_TYPE.CHARACTER)
 hs.font.color.rgb=RGBColor.from_string('3B5963');hs.font.underline=True
 h=sec.header.paragraphs[0];h.style='Header';h.text='Елена Дым  ·  Как продавать услуги'
 f=sec.footer.paragraphs[0];f.style='Footer';f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 r=f.add_run();r._r.append(xml('fldChar',fldCharType='begin'));r=f.add_run();it=xml('instrText');it.text=' PAGE ';r._r.append(it);r=f.add_run();r._r.append(xml('fldChar',fldCharType='end'))
 for obj in [sec.first_page_header,sec.first_page_footer]:obj.paragraphs[0].text=''
 doc.core_properties.title='Как продавать услуги';doc.core_properties.author='Елена Дым';doc.core_properties.subject='Секреты продвижения услуг. Книга первая';doc.core_properties.language='ru-RU';doc.core_properties.comments=''
 doc.add_paragraph('СЕКРЕТЫ ПРОДВИЖЕНИЯ УСЛУГ',style='Cover Series')
 doc.add_paragraph('Елена Дым',style='Cover Author')
 doc.add_paragraph('Как продавать\nуслуги',style='Title')
 doc.add_paragraph('Книга первая',style='Subtitle');doc.add_paragraph('2026',style='Cover Year')
 doc.add_paragraph('Содержание',style='Contents Title')
 for ch,p in files:
  if ch in PARTS:doc.add_paragraph(PARTS[ch],style='Contents Part')
  title=p.read_text().splitlines()[0].lstrip('# ')
  para=doc.add_paragraph(style='Contents Entry');link(para,title,anchor='chapter_'+str(ch))
 records=[];bid=num_id(doc)
 for ch,p in files:body(doc,p.read_text(),ch,records,bid)
 for t in doc.tables:
  pr=t._tbl.tblPr;assert pr.find(qn('w:tblW')).get(qn('w:w'))=='9360';assert pr.find(qn('w:tblInd')).get(qn('w:w'))=='120'
  ws=[int(c.get(qn('w:w'))) for c in t._tbl.tblGrid];assert sum(ws)==9360
  for row in t.rows:assert [int(c._tc.tcPr.find(qn('w:tcW')).get(qn('w:w'))) for c in row.cells]==ws
 doc.save(out)
 return records,{'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'bibliography_entries':sum(r['kind']=='bibliography' for r in records),'source_blocks':len(records)}

def slug(s):return re.sub(r'[^\w\s-]','',s.lower()).replace(' ','-')
def make_md(files,out):
 chunks=['# Как продавать услуги\n\nЕлена Дым\n\nСерия «Секреты продвижения услуг». Книга первая.\n\n2026\n\n## Содержание\n']
 for ch,p in files:
  title=p.read_text().splitlines()[0].lstrip('# ')
  if ch in PARTS:chunks.append('\n\n**'+PARTS[ch]+'**\n')
  chunks.append('\n- ['+title+'](#'+slug(title)+')')
 chunks.append('\n\n')
 for ch,p in files:
  if ch in PARTS:chunks.append('## '+PARTS[ch]+'\n\n')
  text=p.read_text()
  add=2 if 1<=ch<=14 else 1
  text=re.sub(r'^(#+) ',lambda m:'#'*add+m.group(1)+' ',text,flags=re.M)
  chunks.append(text.rstrip()+'\n\n')
 out.write_text(''.join(chunks))

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--output-dir',required=True);args=ap.parse_args();out=Path(args.output_dir).resolve();out.mkdir(parents=True,exist_ok=True)
 files=[(int(p.name[:2]),p) for p in sorted((BASE/'manuscript').glob('*.md'))]
 assert [n for n,p in files]==list(range(16))+[99]
 dp=out/'Как-продавать-услуги_Полная-рукопись.docx';mp=out/'Как-продавать-услуги_Полная-рукопись.md'
 records,counts=make_doc(files,dp);make_md(files,mp)
 audit={'tokens':TOKENS,'counts':counts,'sources':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for n,p in files},'records':records}
 (out/'build-audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2))
 print(json.dumps({'docx':str(dp),'markdown':str(mp),**counts},ensure_ascii=False))

if __name__=='__main__':main()
