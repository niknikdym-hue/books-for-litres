"""Prepare the matching Litres DOCX from the newly built reading edition.

Primary-runtime dependencies: python-docx, reportlab, PyMuPDF, Pillow.
Rules checked 2026-09-07: https://selfpub.ru/faq/first-step/
and https://blog.selfpub.ru/workinword . This does not publish the book.
Tables and formulas become exact, embedded PNGs; their text is retained in alt
descriptions and the audit. Chapter prose and bibliographic entries are preserved.
"""
from pathlib import Path
import argparse, hashlib, json, re
from io import BytesIO
from zipfile import ZipFile
from PIL import Image
from xml.sax.saxutils import escape
from urllib.parse import urlparse
import fitz
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph as WordParagraph
from docx.table import Table as WordTable
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.pdfgen.canvas import Canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, Table, TableStyle
from reportlab.lib import colors
from build_book import widths

ROOT=Path(__file__).resolve().parent
FONT_ROOT=Path('/usr/share/fonts/truetype/dejavu')
for name,file in [('TableFont','DejaVuSans.ttf'),('TableBold','DejaVuSans-Bold.ttf'),
                  ('FormulaFont','DejaVuSerif.ttf'),('FormulaBold','DejaVuSerif-Bold.ttf')]:
    pdfmetrics.registerFont(TTFont(name,str(FONT_ROOT/file)))

# Remove commercial-site addresses while keeping source authors, titles, dates,
# and the link labels as plain text. Government, academic, DOI and professional
# association references retain their non-promotional source links.
COMMERCIAL_DOMAINS={'2bobs.com','punctuation.com','qlean.ru','robfitz.com',
 'ru.jura.com','saasclub.io','therewiredgroup.com','trustedadvisor.com',
 'www.academia.edu','www.consultant.ru','link.springer.com'}

FORMULAS=[
 '6 000 ₽ + 12 часов × 700 ₽ + 12 часов × 100 ₽ = 15 600 ₽.',
 'Остаток = 0,96 × Ц − С.',
 '15 600 ₽ ÷ 0,96 = 16 250 ₽.',
 '(15 600 ₽ + 3 000 ₽) ÷ 0,96 = 19 375 ₽.',
 '0,04 × Ц','19 700 ₽ ÷ 0,96',
 '25 + 3 × 70 + 2 × 40 + 30 = 345',
 '25 + 4 × 70 + 3 × 40 + 30 = 455',
 '110 × n + 15 ≤ 360','110 × n + 15',
]
MATH_PATTERN=re.compile('|'.join(re.escape(x) for x in sorted(FORMULAS,key=len,reverse=True)))

def text_of(el):
    return ''.join(x.text or '' for x in el.iter(qn('w:t')))

def save_illustration(pdf_bytes,path):
    """Rasterize exact vector content within Litres' recommended pixel bounds.

    Pixel dimensions belong to the embedded original, not its Word display size.
    72 dpi is metadata; explicit Word dimensions preserve the intended layout.
    """
    with fitz.open(stream=pdf_bytes,filetype='pdf') as pdf:
        page=pdf[0]
        scale=min(300/72,1024/page.rect.width,1024/page.rect.height)
        pix=page.get_pixmap(matrix=fitz.Matrix(scale,scale),colorspace=fitz.csRGB,alpha=False)
        assert pix.width<=1024 and pix.height<=1024
        assert pix.width*pix.height<=2_000_000
        pix.set_dpi(72,72)
        pix.save(str(path))

def make_png(flow,width,height,path):
    buf=BytesIO(); c=Canvas(buf,pagesize=(width,height))
    c.setFillColor(colors.white);c.rect(0,0,width,height,stroke=0,fill=1)
    flow.drawOn(c,0,height-flow.wrap(width,height)[1]);c.showPage();c.save()
    save_illustration(buf.getvalue(),path)

def table_png(rows,ch,idx,path):
    colwidths=[w/20 for w in widths(ch,idx,len(rows[0]))]
    normal=ParagraphStyle('table',fontName='TableFont',fontSize=9.5,leading=12,splitLongWords=False)
    head=ParagraphStyle('head',parent=normal,fontName='TableBold',fontSize=9.3)
    # Explicit discretionary hyphens for the three words wider than narrow cells.
    # Source strings and alt descriptions retain the original, unbroken words.
    hyphens={'самостоятельных':'самостоятель\u00adных',
             'предоставленном':'предоставлен\u00adном',
             'Предоставленные':'Предоставлен\u00adные'}
    cells=[]
    for ri,row in enumerate(rows):
        line=[];style=head if ri==0 else normal
        for ci,s in enumerate(row):
            for original,hyphenated in hyphens.items():s=s.replace(original,hyphenated)
            for word in s.split():
                chunks=word.split('\u00ad')
                for j,chunk in enumerate(chunks):
                    chunk+= '-' if j<len(chunks)-1 else ''
                    assert pdfmetrics.stringWidth(chunk,style.fontName,style.fontSize)<=colwidths[ci]-12,('Word too wide',ch,idx,ci,word)
            line.append(Paragraph(escape(s),style))
        cells.append(line)
    table=Table(cells,colWidths=colwidths)
    table.setStyle(TableStyle([
      ('GRID',(0,0),(-1,-1),.4,colors.HexColor('#B6B6B6')),
      ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#F3F4F5')),
      ('VALIGN',(0,0),(-1,-1),'TOP'),
      ('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),
      ('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),
    ]))
    w,h=table.wrap(468,10000)
    assert h<640,('Table taller than usable page',ch,idx,h)
    make_png(table,w,h,path)
    return w,h

def formula_png(s,bold,path):
    font='FormulaBold' if bold else 'FormulaFont'
    size=12
    width=pdfmetrics.stringWidth(s,font,size)+2
    assert width<=468,(s,width)
    buf=BytesIO();c=Canvas(buf,pagesize=(width,16))
    c.setFillColor(colors.white);c.rect(0,0,width,16,stroke=0,fill=1)
    c.setFillColor(colors.black);c.setFont(font,size);c.drawString(1,3,s)
    c.showPage();c.save()
    save_illustration(buf.getvalue(),path)
    return width,16

def set_alt(shape,title,description):
    for side in ('T','B','L','R'):shape._inline.set('dist'+side,'0')
    shape._inline.docPr.set('title',title)
    shape._inline.docPr.set('descr',description)

def semantic_blocks(doc):
    blocks=[]
    for e in doc.element.body:
        if e.tag==qn('w:p'):
            s=''
            for node in e.iter():
                if node.tag==qn('w:t'):s+=node.text or ''
                elif node.tag.endswith('}docPr'):
                    alt=node.get('descr','')
                    if node.get('title','').startswith('Таблица'):
                        rows=json.loads(alt)
                        for row in rows:
                            for cell in row:blocks.append(cell)
                    else:s+=alt
            if s:blocks.append(s)
        elif e.tag==qn('w:tbl'):
            for row in WordTable(e,doc).rows:
                for cell in row.cells:blocks.append(cell.text)
    return blocks

def validate_structure(doc):
    ns={'wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
        'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    root=doc.element
    count=lambda tag:len(list(root.iter(qn(tag))))
    result={'inline_images':len(list(root.iter('{'+ns['wp']+'}inline'))),
      'floating_images':len(list(root.iter('{'+ns['wp']+'}anchor'))),
      'linked_images':len([e for e in root.iter('{'+ns['a']+'}blip') if e.get(qn('r:link'))]),
      'image_hyperlinks':len(list(root.iter('{'+ns['a']+'}hlinkClick'))),
      'header_footer_references':count('w:headerReference')+count('w:footerReference'),
      'page_breaks':len([e for e in root.iter(qn('w:br')) if e.get(qn('w:type'))=='page']),
      'page_break_before':count('w:pageBreakBefore'),
      'native_tables':len(doc.tables),
      'top_headings':len([p for p in doc.paragraphs if p.style.name=='Heading 1'])}
    assert result=={'inline_images':26,'floating_images':0,'linked_images':0,
      'image_hyperlinks':0,'header_footer_references':0,'page_breaks':0,
      'page_break_before':0,'native_tables':0,'top_headings':17},result
    return result

def main(argv=None):
    ap=argparse.ArgumentParser();ap.add_argument('--reading-docx',required=True)
    ap.add_argument('--output-dir',required=True);args=ap.parse_args(argv)
    source=Path(args.reading_docx).resolve();out=Path(args.output_dir).resolve()
    out.mkdir(parents=True,exist_ok=True);assets=out/'litres-assets';assets.mkdir(exist_ok=True)
    doc=Document(source)
    # Remove cover and reading TOC, stopping at the actual introduction heading.
    first=next(p._p for p in doc.paragraphs if p.style.name=='Heading 1')
    for e in list(doc.element.body):
        if e is first:break
        doc.element.body.remove(e)
    before=semantic_blocks(doc)
    removed_links=[]
    for h in list(doc.element.body.iter(qn('w:hyperlink'))):
        rid=h.get(qn('r:id'));rel=doc.part.rels.get(rid) if rid else None
        if rel is not None and urlparse(rel.target_ref).netloc in COMMERCIAL_DOMAINS:
            removed_links.append({'label':text_of(h),'url':rel.target_ref})
            parent=h.getparent();pos=parent.index(h)
            for child in list(h):
                parent.insert(pos,child);pos+=1
                for rs in list(child.iter(qn('w:rStyle'))):rs.getparent().remove(rs)
            parent.remove(h)
    # Remove references to page furniture and the now-unused package parts.
    for rid,rel in list(doc.part.rels.items()):
        if rel.reltype in (RT.HEADER,RT.FOOTER):doc.part.drop_rel(rid)
        elif rel.reltype==RT.HYPERLINK:
            if not any(h.get(qn('r:id'))==rid for h in doc.element.body.iter(qn('w:hyperlink'))):
                doc.part.drop_rel(rid)
    for root in (doc.element,doc.styles.element):
        for tag in ('pageBreakBefore','lastRenderedPageBreak','headerReference','footerReference','titlePg'):
            for e in list(root.iter(qn('w:'+tag))):e.getparent().remove(e)
        for e in list(root.iter(qn('w:br'))):
            if e.get(qn('w:type'))=='page':e.getparent().remove(e)
    for k in ('title','author','subject','comments','keywords'):
        setattr(doc.core_properties,k,'')
    for n in ('Heading 1','Heading 2','Heading 3'):
        doc.styles[n].font.bold=False;doc.styles[n].font.italic=False
        doc.styles[n].font.color.rgb=RGBColor(0,0,0)
    # Converter receives ordinary paragraphs, genuine heading styles and numbering.
    for p in doc.paragraphs:
        if p.style.name in ('Book Excerpt','Bibliography Text','Table After'):p.style='Normal'
        if p.style.name.startswith('Heading'):
            for r in p.runs:r.bold=None;r.italic=None
        for e in list(p._p.iter(qn('w:bookmarkStart')))+list(p._p.iter(qn('w:bookmarkEnd'))):
            e.getparent().remove(e)
    table_assets=[];formula_assets=[];ch=0;idx=0;mathidx=0
    for e in list(doc.element.body):
        if e.tag==qn('w:p'):
            p=WordParagraph(e,doc)
            if p.style.name=='Heading 1':
                m=re.match(r'Глава (\d+)\.',p.text)
                ch=int(m.group(1)) if m else 0;idx=0
            s=p.text
            matches=list(MATH_PATTERN.finditer(s))
            if not matches:continue
            assert not list(p._p.iter(qn('w:hyperlink')))
            for child in list(p._p):
                if child.tag!=qn('w:pPr'):p._p.remove(child)
            pos=0
            for m in matches:
                if m.start()>pos:p.add_run(s[pos:m.start()])
                formula=m.group();mathidx+=1;path=assets/f'formula-{mathidx:02}.png'
                w,h=formula_png(formula,len(matches)==1 and formula==s,path)
                run=p.add_run();shape=run.add_picture(str(path),width=Pt(w),height=Pt(h))
                # A slight baseline adjustment aligns inline raster math with prose.
                if formula!=s:
                    from build_book import xml
                    run._r.get_or_add_rPr().append(xml('position',val='-5'))
                set_alt(shape,'Формула '+str(mathidx),formula)
                formula_assets.append({'file':path.name,'text':formula,'chapter':ch,'width_pt':w,'height_pt':h})
                pos=m.end()
            if pos<len(s):p.add_run(s[pos:])
        elif e.tag==qn('w:tbl'):
            table=WordTable(e,doc);rows=[[c.text for c in r.cells] for r in table.rows]
            path=assets/f'table-{ch:02}-{idx+1:02}.png';w,h=table_png(rows,ch,idx,path)
            p=doc.add_paragraph();p.paragraph_format.keep_together=True
            shape=p.add_run().add_picture(str(path),width=Inches(6.5),height=Pt(h))
            set_alt(shape,f'Таблица. Глава {ch}, {idx+1}',json.dumps(rows,ensure_ascii=False))
            e.addprevious(p._p);e.getparent().remove(e)
            table_assets.append({'file':path.name,'chapter':ch,'index':idx,'rows':rows,'width_pt':w,'height_pt':h})
            idx+=1
    for p in list(doc.paragraphs):
        if not p.text.strip() and not list(p._p.iter(qn('w:drawing'))):p._p.getparent().remove(p._p)
    after=semantic_blocks(doc);assert before==after,'Content differs after conversion'
    assert len(table_assets)==16 and len(formula_assets)==10
    assert len([p for p in doc.paragraphs if p.style.name=='Heading 1'])==17
    assert not doc.tables
    assert not any(p.text.startswith('Часть ') for p in doc.paragraphs)
    forbidden='ˊˈʻʼˋʹːˌ'
    assert not any(c in ''.join(after) for c in forbidden)
    target=out/'Как-продавать-услуги_Для-Литрес.docx';doc.save(target)
    assert target.stat().st_size<=70_000_000
    # Verify the saved file rather than trusting the in-memory transformation.
    reopened=Document(target);assert semantic_blocks(reopened)==before
    structure_checks=validate_structure(reopened)
    image_checks=[]
    with ZipFile(target) as archive:
        for name in archive.namelist():
            if not name.startswith('word/media/'):continue
            with Image.open(BytesIO(archive.read(name))) as im:
                dpi=im.info.get('dpi')
                assert im.format=='PNG' and im.mode=='RGB',(name,im.format,im.mode)
                assert im.width<=1024 and im.height<=1024,(name,im.size)
                assert im.width*im.height<=2_000_000,(name,im.size)
                assert dpi and all(abs(x-72)<.1 for x in dpi),(name,dpi)
                image_checks.append({'file':name,'width_px':im.width,'height_px':im.height,
                  'pixels':im.width*im.height,'format':im.format,'mode':im.mode,'dpi':dpi})
    assert len(image_checks)==26
    audit={'rules_checked':'2026-09-07','reading_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),
      'litres_sha256':hashlib.sha256(target.read_bytes()).hexdigest(),'size_bytes':target.stat().st_size,
      'semantic_blocks':len(after),'text_preserved':True,'table_images':table_assets,
      'formula_images':formula_assets,'commercial_links_as_plain_citations':removed_links,
      'chapter_headings':17,'native_tables':0,'heading_parts':0,
      'embedded_image_checks':image_checks,'image_checks_passed':True,
      'structure_checks':structure_checks,
      'image_rules':{'max_width_px':1024,'max_height_px':1024,'max_pixels':2_000_000,
                     'dpi':72,'format':'PNG','mode':'RGB','wrapping':'inline'},
      'platform_conversion_tested':False}
    (out/'litres-audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2))
    print(json.dumps({k:v for k,v in audit.items() if k not in ('table_images','formula_images','commercial_links_as_plain_citations','embedded_image_checks')},ensure_ascii=False))

if __name__=='__main__':main()
