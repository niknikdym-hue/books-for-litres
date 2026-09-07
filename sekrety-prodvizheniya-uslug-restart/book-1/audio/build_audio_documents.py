"""Create the two requested DOCX editions and a clean narration TXT."""
from pathlib import Path
import sys, argparse, json, hashlib, shutil
from urllib.parse import urlparse
from zipfile import ZipFile
from docx import Document
from docx.shared import RGBColor, Pt
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

HERE=Path(__file__).resolve().parent
sys.path.insert(0,str(HERE.parent))
import build_book as book
from build_litres import semantic_blocks, text_of, COMMERCIAL_DOMAINS

def make_litres(source,target):
    doc=Document(source)
    first=next(p._p for p in doc.paragraphs if p.style.name=='Heading 1')
    for e in list(doc.element.body):
        if e is first:break
        doc.element.body.remove(e)
    before=semantic_blocks(doc)
    removed=[]
    for h in list(doc.element.body.iter(qn('w:hyperlink'))):
        rid=h.get(qn('r:id'));rel=doc.part.rels.get(rid) if rid else None
        if rel and urlparse(rel.target_ref).netloc in COMMERCIAL_DOMAINS:
            removed.append({'label':text_of(h),'url':rel.target_ref})
            parent=h.getparent();pos=parent.index(h)
            for child in list(h):
                parent.insert(pos,child);pos+=1
                for rs in list(child.iter(qn('w:rStyle'))):rs.getparent().remove(rs)
            parent.remove(h)
    for rid,rel in list(doc.part.rels.items()):
        if rel.reltype in (RT.HEADER,RT.FOOTER):doc.part.drop_rel(rid)
        elif rel.reltype==RT.HYPERLINK and not any(h.get(qn('r:id'))==rid for h in doc.element.body.iter(qn('w:hyperlink'))):doc.part.drop_rel(rid)
    for root in (doc.element,doc.styles.element):
        for tag in ('pageBreakBefore','lastRenderedPageBreak','headerReference','footerReference','titlePg'):
            for e in list(root.iter(qn('w:'+tag))):e.getparent().remove(e)
        for e in list(root.iter(qn('w:br'))):
            if e.get(qn('w:type'))=='page':e.getparent().remove(e)
    for k in ('title','author','subject','comments','keywords'):setattr(doc.core_properties,k,'')
    for n in ('Heading 1','Heading 2','Heading 3'):
        doc.styles[n].font.bold=False;doc.styles[n].font.italic=False;doc.styles[n].font.color.rgb=RGBColor(0,0,0)
    for p in list(doc.paragraphs):
        if p.style.name in ('Book Excerpt','Bibliography Text','Table After'):p.style='Normal'
        if p.style.name.startswith('Heading'):
            for r in p.runs:r.bold=None;r.italic=None
        for tag in ('bookmarkStart','bookmarkEnd'):
            for e in list(p._p.iter(qn('w:'+tag))):e.getparent().remove(e)
        if not p.text.strip():p._p.getparent().remove(p._p)
    assert semantic_blocks(doc)==before
    doc.save(target)
    reread=Document(target)
    assert semantic_blocks(reread)==before
    root=reread.element
    result={
      'headings':sum(p.style.name=='Heading 1' for p in reread.paragraphs),
      'tables':len(reread.tables),'images':len(reread.inline_shapes),
      'page_breaks':sum(e.get(qn('w:type'))=='page' for e in root.iter(qn('w:br'))),
      'page_break_before':len(list(root.iter(qn('w:pageBreakBefore')))),
      'furniture':len(list(root.iter(qn('w:headerReference'))))+len(list(root.iter(qn('w:footerReference')))),
      'semantic_blocks':len(before),'text_preserved':True,
      'size_bytes':target.stat().st_size,'commercial_links_plain':len(removed),
      'platform_conversion_tested':False}
    assert result['headings']==17
    assert all(result[x]==0 for x in ('tables','images','page_breaks','page_break_before','furniture'))
    assert target.stat().st_size<70_000_000
    assert not any(c in ''.join(before) for c in 'ˊˈʻʼˋʹːˌ')
    with ZipFile(target) as z:assert not any(p.startswith('word/media/') for p in z.namelist())
    return result

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--output-dir',required=True);args=ap.parse_args()
    out=Path(args.output_dir).resolve();out.mkdir(parents=True,exist_ok=True)
    files=[(int(p.name[:2]),p) for p in sorted((HERE/'manuscript').glob('*.md'))]
    assert [n for n,p in files]==list(range(16))+[99]
    reading=out/'Как-продавать-услуги_Аудиоредакция_Для-чтения.docx'
    litres=out/'Как-продавать-услуги_Аудиоредакция_Для-Литрес.docx'
    records,counts=book.make_doc(files,reading)
    doc=Document(reading)
    # Named reading-only fit overrides prevent one-paragraph chapter tails.
    chapter=-1
    for p in doc.paragraphs:
        if p.style.name=='Subtitle':p.text='Книга первая. Аудиоредакция'
        if p.style.name=='Heading 1':chapter+=1
        if p.style.name=='Normal' and chapter in (0,2,4):
            p.paragraph_format.space_after=Pt(2)
            if chapter==0:p.paragraph_format.line_spacing=1.10
    doc.core_properties.subject='Секреты продвижения услуг. Книга первая. Аудиоредакция 1.0'
    for section in doc.sections:
        section.header.paragraphs[0].text='Елена Дым  ·  Как продавать услуги  ·  Аудиоредакция'
    doc.save(reading)
    assert not doc.tables and not doc.inline_shapes
    lit=make_litres(reading,litres)
    txt=out/'Как-продавать-услуги_Аудиоредакция_Для-озвучки.txt'
    shutil.copyfile(HERE/'Как-продавать-услуги_Аудиоредакция.txt',txt)
    audit={'edition':'audio-1.0','source_commit':'e94c055514acc7f2643265c03494cd5677e056d6',
      'reading':counts,'litres':lit,'files':{p.name:{'size_bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in (reading,litres,txt)},
      'spoken_chapters':16,'bibliography_entries':38,'bibliography_spoken':False,
      'english_examples_in_text_appendix':True,'audio_recorded':False,'listen_review_done':False,
      'tokens':book.TOKENS,
      'reading_fit_overrides':{'introduction':{'paragraph_after_pt':2,'line_spacing':1.10},'chapters_2_and_4':{'paragraph_after_pt':2},'reason':'Avoid almost-empty final pages without changing text, font size, or chapter opening rules.'}}
    (out/'audio-build-audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(audit,ensure_ascii=False))

if __name__=='__main__':main()
